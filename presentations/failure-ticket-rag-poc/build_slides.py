#!/usr/bin/env python3
"""障害票RAG PoC 発表資料（16枚）を PowerPoint / HTML プレビューで生成する。"""

from __future__ import annotations

from pathlib import Path

from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_SHAPE
from pptx.enum.text import PP_ALIGN, MSO_ANCHOR
from pptx.util import Inches, Pt, Emu

OUT_DIR = Path(__file__).resolve().parent
PPTX_PATH = OUT_DIR / "障害票RAG_PoC発表資料.pptx"
HTML_PATH = OUT_DIR / "preview.html"

# スライドサイズ: ワイド 16:9
SLIDE_W = Inches(13.333)
SLIDE_H = Inches(7.5)

NAVY = RGBColor(0x1B, 0x3A, 0x5F)
ACCENT = RGBColor(0x2E, 0x6B, 0x9E)
LIGHT_BG = RGBColor(0xF5, 0xF7, 0xFA)
WHITE = RGBColor(0xFF, 0xFF, 0xFF)
DARK = RGBColor(0x2C, 0x2C, 0x2C)
GRAY = RGBColor(0x5A, 0x5A, 0x5A)
RED_SOFT = RGBColor(0xA6, 0x3D, 0x40)
GREEN_SOFT = RGBColor(0x2D, 0x6A, 0x4F)
ORANGE = RGBColor(0xC4, 0x6B, 0x2E)

FONT = "Yu Gothic"
FONT_FALLBACK = "MS Gothic"

EXPECTED_TITLES = [
    "タイトル",
    "背景・目的",
    "システム全体構成",
    "Basic RAG",
    "Basic RAGで発生した課題",
    "Query Rewrite RAG",
    "Query Rewriteによる改善結果",
    "Critic RAG",
    "Critic改善",
    "Self-RAG",
    "Self-RAG改善",
    "試行錯誤の変遷",
    "評価結果",
    "PoCで得られた知見",
    "今後の展望",
    "まとめ",
]


def set_run_font(run, size_pt: int, bold: bool = False, color: RGBColor = DARK, font_name: str = FONT):
    run.font.name = font_name
    run.font.size = Pt(size_pt)
    run.font.bold = bold
    run.font.color.rgb = color
    # East Asian font hint for Japanese
    rPr = run._r.get_or_add_rPr()
    from pptx.oxml.ns import qn
    from lxml import etree

    ea = rPr.find(qn("a:ea"))
    if ea is None:
        ea = etree.SubElement(rPr, qn("a:ea"))
    ea.set("typeface", font_name)


def add_textbox(slide, left, top, width, height, text, size=18, bold=False, color=DARK, align=PP_ALIGN.LEFT):
    shape = slide.shapes.add_textbox(left, top, width, height)
    tf = shape.text_frame
    tf.word_wrap = True
    p = tf.paragraphs[0]
    p.alignment = align
    run = p.add_run()
    run.text = text
    set_run_font(run, size, bold=bold, color=color)
    return shape


def add_bullets(slide, left, top, width, height, items: list[str], size=16, color=DARK, spacing=Pt(8)):
    shape = slide.shapes.add_textbox(left, top, width, height)
    tf = shape.text_frame
    tf.word_wrap = True
    for i, item in enumerate(items):
        p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        p.alignment = PP_ALIGN.LEFT
        p.space_after = spacing
        run = p.add_run()
        run.text = f"• {item}"
        set_run_font(run, size, color=color)
    return shape


def add_rect(slide, left, top, width, height, fill: RGBColor, line=None):
    shape = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, left, top, width, height)
    shape.fill.solid()
    shape.fill.fore_color.rgb = fill
    if line is None:
        shape.line.fill.background()
    else:
        shape.line.color.rgb = line
    return shape


def add_rounded_box(slide, left, top, width, height, fill: RGBColor, text: str, size=14, bold=True, text_color=WHITE):
    shape = slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, left, top, width, height)
    shape.fill.solid()
    shape.fill.fore_color.rgb = fill
    shape.line.fill.background()
    tf = shape.text_frame
    tf.word_wrap = True
    tf.paragraphs[0].alignment = PP_ALIGN.CENTER
    shape.text_frame.auto_size = None
    # vertical center via anchor
    try:
        tf._txBody.bodyPr.set("anchor", "ctr")
    except Exception:
        pass
    p = tf.paragraphs[0]
    run = p.add_run()
    run.text = text
    set_run_font(run, size, bold=bold, color=text_color)
    return shape


def add_header_bar(slide, title: str, subtitle: str | None = None):
    add_rect(slide, Inches(0), Inches(0), SLIDE_W, Inches(1.05), NAVY)
    add_textbox(slide, Inches(0.5), Inches(0.22), Inches(12), Inches(0.55), title, size=26, bold=True, color=WHITE)
    if subtitle:
        add_textbox(slide, Inches(0.5), Inches(0.68), Inches(12), Inches(0.3), subtitle, size=12, color=RGBColor(0xC8, 0xD6, 0xE5))
    # accent line
    add_rect(slide, Inches(0), Inches(1.05), SLIDE_W, Inches(0.06), ACCENT)


def add_footer(slide, page: int, total: int = 16):
    add_textbox(
        slide,
        Inches(0.5),
        Inches(7.1),
        Inches(10),
        Inches(0.3),
        "障害票RAG PoC｜AI技術検証プロジェクト",
        size=10,
        color=GRAY,
    )
    add_textbox(
        slide,
        Inches(11.5),
        Inches(7.1),
        Inches(1.5),
        Inches(0.3),
        f"{page} / {total}",
        size=10,
        color=GRAY,
        align=PP_ALIGN.RIGHT,
    )


def flow_boxes(slide, labels: list[str], top=Inches(2.0), box_w=Inches(2.2), box_h=Inches(0.7), gap=Inches(0.35)):
    n = len(labels)
    total_w = n * box_w + (n - 1) * gap
    left0 = (SLIDE_W - total_w) / 2
    for i, label in enumerate(labels):
        left = left0 + i * (box_w + gap)
        add_rounded_box(slide, left, top, box_w, box_h, ACCENT if i % 2 == 0 else NAVY, label, size=13)
        if i < n - 1:
            arrow_left = left + box_w + Inches(0.02)
            add_textbox(slide, arrow_left, top + Inches(0.1), gap, Inches(0.5), "→", size=22, bold=True, color=NAVY, align=PP_ALIGN.CENTER)


# ---------------------------------------------------------------------------
# Slide content builders
# ---------------------------------------------------------------------------

def slide_01_title(prs):
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    add_rect(slide, Inches(0), Inches(0), SLIDE_W, SLIDE_H, NAVY)
    add_rect(slide, Inches(0), Inches(5.8), SLIDE_W, Inches(1.7), ACCENT)
    add_textbox(slide, Inches(0.8), Inches(1.8), Inches(11.5), Inches(0.4), "AI技術検証プロジェクト｜PoC報告", size=16, color=RGBColor(0xC8, 0xD6, 0xE5))
    add_textbox(
        slide,
        Inches(0.8),
        Inches(2.4),
        Inches(11.5),
        Inches(1.5),
        "障害票を活用した\n運用改善テーマ発見AIのPoC検証",
        size=36,
        bold=True,
        color=WHITE,
    )
    add_textbox(
        slide,
        Inches(0.8),
        Inches(4.3),
        Inches(11.5),
        Inches(0.5),
        "Basic RAG → Query Rewrite → Critic → Self-RAG による段階的改善",
        size=16,
        color=RGBColor(0xC8, 0xD6, 0xE5),
    )
    add_textbox(slide, Inches(0.8), Inches(6.2), Inches(11.5), Inches(0.4), "ドキュメント作成担当 ／ AI技術検証プロジェクト", size=14, color=WHITE)
    return "タイトル"


def slide_02_background(prs):
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    add_header_bar(slide, "背景・目的", "発表のゴール")
    add_bullets(
        slide,
        Inches(0.7),
        Inches(1.5),
        Inches(11.5),
        Inches(2.2),
        [
            "障害票をナレッジとして蓄積し、運用改善テーマを自動生成するAIを構築した",
            "本発表では「システム紹介」ではなく、構築〜改善の試行錯誤プロセスを説明する",
            "各段階を「発生した課題 → 原因分析 → 改善 → 結果」の流れで整理する",
        ],
        size=18,
    )
    # narrative box
    add_rect(slide, Inches(0.7), Inches(4.0), Inches(11.8), Inches(2.4), LIGHT_BG)
    add_textbox(slide, Inches(1.0), Inches(4.2), Inches(11), Inches(0.4), "本資料の語り方", size=16, bold=True, color=NAVY)
    flow_boxes(slide, ["発生した課題", "原因分析", "改善", "結果"], top=Inches(4.9), box_w=Inches(2.4), box_h=Inches(0.85))
    add_footer(slide, 2)
    return "背景・目的"


def slide_03_architecture(prs):
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    add_header_bar(slide, "システム全体構成", "Dify 上に構築した障害票ナレッジ活用AI")
    add_bullets(
        slide,
        Inches(0.7),
        Inches(1.4),
        Inches(11.5),
        Inches(1.4),
        [
            "障害票をナレッジベースとして登録（Dify）",
            "ユーザーが入力した障害内容に対し、以下3機能を実行",
        ],
        size=17,
    )
    caps = [
        ("類似障害検索", "過去の類似チケットを検索"),
        ("障害情報抽出", "原因・対応・恒久対策を整理"),
        ("改善テーマ生成", "テーマ・内容・期待効果を提案"),
    ]
    for i, (title, desc) in enumerate(caps):
        left = Inches(0.7) + i * Inches(4.0)
        add_rounded_box(slide, left, Inches(3.1), Inches(3.6), Inches(0.7), ACCENT, title, size=16)
        add_textbox(slide, left, Inches(4.0), Inches(3.6), Inches(0.8), desc, size=14, color=GRAY, align=PP_ALIGN.CENTER)

    add_textbox(slide, Inches(0.7), Inches(5.2), Inches(11.5), Inches(0.4), "技術進化の軸", size=15, bold=True, color=NAVY)
    flow_boxes(slide, ["Basic RAG", "Query Rewrite", "Critic RAG", "Self-RAG"], top=Inches(5.7), box_w=Inches(2.5), box_h=Inches(0.65), gap=Inches(0.3))
    add_footer(slide, 3)
    return "システム全体構成"


def slide_04_basic_rag(prs):
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    add_header_bar(slide, "Basic RAG", "初期構成：検索 → 抽出 → テーマ生成")
    flow_boxes(
        slide,
        ["入力", "ナレッジ検索", "障害情報抽出", "改善テーマ生成"],
        top=Inches(1.4),
        box_w=Inches(2.5),
        box_h=Inches(0.7),
        gap=Inches(0.3),
    )

    # left prompt card
    add_rect(slide, Inches(0.5), Inches(2.5), Inches(6.0), Inches(4.2), LIGHT_BG)
    add_textbox(slide, Inches(0.7), Inches(2.65), Inches(5.6), Inches(0.35), "【障害情報抽出】使用プロンプト", size=14, bold=True, color=NAVY)
    add_bullets(
        slide,
        Inches(0.7),
        Inches(3.15),
        Inches(5.6),
        Inches(3.3),
        [
            "あなたは障害分析担当です。",
            "検索結果から次を整理する：",
            "障害名 / 原因 / 対応 / 恒久対策",
        ],
        size=15,
    )

    # right prompt card
    add_rect(slide, Inches(6.8), Inches(2.5), Inches(6.0), Inches(4.2), LIGHT_BG)
    add_textbox(slide, Inches(7.0), Inches(2.65), Inches(5.6), Inches(0.35), "【改善テーマ生成】使用プロンプト", size=14, bold=True, color=NAVY)
    add_bullets(
        slide,
        Inches(7.0),
        Inches(3.15),
        Inches(5.6),
        Inches(3.3),
        [
            "あなたは運用改善担当です。",
            "障害原因を分析し、次を作成する：",
            "改善テーマ / 改善内容 / 期待効果",
        ],
        size=15,
    )
    add_footer(slide, 4)
    return "Basic RAG"


def slide_05_basic_issues(prs):
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    add_header_bar(slide, "Basic RAGで発生した課題", "課題 → 原因分析 → 学び")

    # success
    add_rounded_box(slide, Inches(0.5), Inches(1.4), Inches(4.0), Inches(0.55), GREEN_SOFT, "成功例", size=14)
    add_bullets(
        slide,
        Inches(0.5),
        Inches(2.1),
        Inches(4.0),
        Inches(2.0),
        ["バックアップ失敗", "DB接続エラー", "容量不足による障害"],
        size=15,
    )

    # issues
    add_rounded_box(slide, Inches(4.8), Inches(1.4), Inches(4.0), Inches(0.55), RED_SOFT, "課題：表現ゆれに弱い", size=14)
    add_bullets(
        slide,
        Inches(4.8),
        Inches(2.1),
        Inches(4.0),
        Inches(2.0),
        ["「DBトラブル」", "「バックアップ問題」", "言い回しが違うと検索が外れる"],
        size=15,
    )

    # learning
    add_rounded_box(slide, Inches(9.1), Inches(1.4), Inches(3.7), Inches(0.55), ORANGE, "学び", size=14)
    add_bullets(
        slide,
        Inches(9.1),
        Inches(2.1),
        Inches(3.7),
        Inches(2.0),
        ["回答品質は検索品質に依存する", "生成改善だけでは限界がある"],
        size=14,
    )

    add_rect(slide, Inches(0.5), Inches(4.6), Inches(12.3), Inches(1.9), LIGHT_BG)
    add_textbox(slide, Inches(0.8), Inches(4.8), Inches(11.8), Inches(0.4), "原因分析", size=16, bold=True, color=NAVY)
    add_bullets(
        slide,
        Inches(0.8),
        Inches(5.3),
        Inches(11.8),
        Inches(1.0),
        [
            "ユーザー発話の表現ゆれを吸収する仕組みが無い",
            "検索クエリが原文のままのため、ナレッジ側の表記と一致しないケースで失敗する",
        ],
        size=15,
    )
    add_footer(slide, 5)
    return "Basic RAGで発生した課題"


def slide_06_query_rewrite(prs):
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    add_header_bar(slide, "Query Rewrite RAG", "目的：表現ゆれ対策")
    add_textbox(slide, Inches(0.7), Inches(1.4), Inches(12), Inches(0.4), "検索前にクエリを書き換えて、ナレッジと突き合わせやすくする", size=16, color=GRAY)

    add_rect(slide, Inches(0.5), Inches(2.0), Inches(12.3), Inches(2.0), LIGHT_BG)
    add_textbox(slide, Inches(0.8), Inches(2.15), Inches(11.8), Inches(0.35), "失敗したプロンプト", size=15, bold=True, color=RED_SOFT)
    add_textbox(
        slide,
        Inches(0.8),
        Inches(2.6),
        Inches(11.8),
        Inches(1.0),
        "「質問を検索しやすい形式へ変換してください」",
        size=18,
        bold=True,
        color=DARK,
    )

    add_rect(slide, Inches(0.5), Inches(4.3), Inches(12.3), Inches(2.2), RGBColor(0xFD, 0xF0, 0xF0))
    add_textbox(slide, Inches(0.8), Inches(4.45), Inches(11.8), Inches(0.35), "問題：変換しすぎた", size=15, bold=True, color=RED_SOFT)
    add_textbox(
        slide,
        Inches(0.8),
        Inches(5.0),
        Inches(11.8),
        Inches(1.0),
        "例）「DBトラブル」 → 「運用改善課題」\n具体的な障害語が消え、検索精度がかえって悪化した",
        size=17,
        color=DARK,
    )
    add_footer(slide, 6)
    return "Query Rewrite RAG"


def slide_07_query_rewrite_result(prs):
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    add_header_bar(slide, "Query Rewriteによる改善結果", "改善 → 結果")

    add_textbox(slide, Inches(0.7), Inches(1.35), Inches(12), Inches(0.35), "改善後プロンプトの重要な制約", size=16, bold=True, color=NAVY)
    constraints = [
        ("原文を維持", "障害語を消さない"),
        ("説明しない", "余計な文章を付けない"),
        ("言い換えすぎない", "過度な一般化を禁止"),
        ("JSONのみ出力", "形式を固定する"),
    ]
    for i, (t, d) in enumerate(constraints):
        left = Inches(0.5) + (i % 4) * Inches(3.15)
        add_rounded_box(slide, left, Inches(1.9), Inches(3.0), Inches(0.6), ACCENT, t, size=14)
        add_textbox(slide, left, Inches(2.65), Inches(3.0), Inches(0.5), d, size=13, color=GRAY, align=PP_ALIGN.CENTER)

    add_textbox(slide, Inches(0.7), Inches(3.5), Inches(12), Inches(0.35), "結果：表現ゆれケースでも検索成功", size=16, bold=True, color=GREEN_SOFT)
    results = [("DBトラブル", "検索成功"), ("バックアップ問題", "検索成功"), ("容量管理不足", "検索成功")]
    for i, (q, r) in enumerate(results):
        left = Inches(0.5) + i * Inches(4.15)
        add_rect(slide, left, Inches(4.1), Inches(3.9), Inches(2.2), LIGHT_BG)
        add_textbox(slide, left + Inches(0.2), Inches(4.4), Inches(3.5), Inches(0.5), q, size=16, bold=True, color=DARK, align=PP_ALIGN.CENTER)
        add_textbox(slide, left + Inches(0.2), Inches(5.0), Inches(3.5), Inches(0.4), "↓", size=18, color=NAVY, align=PP_ALIGN.CENTER)
        add_textbox(slide, left + Inches(0.2), Inches(5.5), Inches(3.5), Inches(0.5), r, size=18, bold=True, color=GREEN_SOFT, align=PP_ALIGN.CENTER)
    add_footer(slide, 7)
    return "Query Rewriteによる改善結果"


def slide_08_critic(prs):
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    add_header_bar(slide, "Critic RAG", "目的：改善テーマの品質評価")
    add_textbox(
        slide,
        Inches(0.7),
        Inches(1.4),
        Inches(12),
        Inches(0.5),
        "検索は改善されたが、「生成された改善テーマ自体の妥当性」を担保する仕組みが必要になった",
        size=15,
        color=GRAY,
    )

    add_rect(slide, Inches(0.5), Inches(2.1), Inches(12.3), Inches(1.8), LIGHT_BG)
    add_textbox(slide, Inches(0.8), Inches(2.25), Inches(11.8), Inches(0.35), "初期プロンプト（失敗の起点）", size=14, bold=True, color=RED_SOFT)
    add_bullets(
        slide,
        Inches(0.8),
        Inches(2.75),
        Inches(11.8),
        Inches(1.0),
        [
            "あなたは運用改善レビュー担当です。",
            "改善テーマの妥当性を評価してください。",
        ],
        size=16,
    )

    add_textbox(slide, Inches(0.7), Inches(4.2), Inches(12), Inches(0.35), "評価軸（導入）", size=15, bold=True, color=NAVY)
    axes = [("関連性", "質問・障害内容とテーマが対応しているか"), ("再発防止効果", "同種障害の再発を抑えられるか"), ("実現可能性", "現場で実行可能な施策か")]
    for i, (t, d) in enumerate(axes):
        left = Inches(0.5) + i * Inches(4.15)
        add_rounded_box(slide, left, Inches(4.7), Inches(3.9), Inches(0.6), ACCENT, t, size=15)
        add_textbox(slide, left, Inches(5.5), Inches(3.9), Inches(0.9), d, size=13, color=GRAY, align=PP_ALIGN.CENTER)
    add_footer(slide, 8)
    return "Critic RAG"


def slide_09_critic_improve(prs):
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    add_header_bar(slide, "Critic改善", "課題 → 原因 → 改善")

    # problem
    add_rect(slide, Inches(0.5), Inches(1.4), Inches(4.0), Inches(3.0), RGBColor(0xFD, 0xF0, 0xF0))
    add_textbox(slide, Inches(0.7), Inches(1.55), Inches(3.6), Inches(0.35), "問題", size=15, bold=True, color=RED_SOFT)
    add_bullets(
        slide,
        Inches(0.7),
        Inches(2.1),
        Inches(3.6),
        Inches(2.0),
        [
            "無関係なテーマでも総合評価◎になるケース",
            "例：ユーザー質問と無関係な提案が高評価",
        ],
        size=14,
    )

    # cause
    add_rect(slide, Inches(4.7), Inches(1.4), Inches(4.0), Inches(3.0), RGBColor(0xFD, 0xF5, 0xE6))
    add_textbox(slide, Inches(4.9), Inches(1.55), Inches(3.6), Inches(0.35), "原因", size=15, bold=True, color=ORANGE)
    add_bullets(
        slide,
        Inches(4.9),
        Inches(2.1),
        Inches(3.6),
        Inches(2.0),
        [
            "ユーザー質問を評価対象に含めていなかった",
            "テーマ単体の「もっともらしさ」だけで採点していた",
        ],
        size=14,
    )

    # improve
    add_rect(slide, Inches(8.9), Inches(1.4), Inches(3.9), Inches(3.0), RGBColor(0xE8, 0xF5, 0xE9))
    add_textbox(slide, Inches(9.1), Inches(1.55), Inches(3.5), Inches(0.35), "改善", size=15, bold=True, color=GREEN_SOFT)
    add_bullets(
        slide,
        Inches(9.1),
        Inches(2.1),
        Inches(3.5),
        Inches(2.0),
        [
            "評価入力にユーザー質問を追加",
            "関連性を必須ゲートにする",
            "低関連は総合評価を落とす",
        ],
        size=14,
    )

    add_rect(slide, Inches(0.5), Inches(4.7), Inches(12.3), Inches(1.8), LIGHT_BG)
    add_textbox(slide, Inches(0.8), Inches(4.9), Inches(11.8), Inches(0.35), "結果", size=15, bold=True, color=NAVY)
    add_bullets(
        slide,
        Inches(0.8),
        Inches(5.4),
        Inches(11.8),
        Inches(0.9),
        [
            "質問と無関係な改善テーマの誤高評価を抑制",
            "関連性・再発防止・実現可能性のバランスでレビュー可能になった",
        ],
        size=15,
    )
    add_footer(slide, 9)
    return "Critic改善"


def slide_10_self_rag(prs):
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    add_header_bar(slide, "Self-RAG", "目的：検索・生成の自己判定と再実行")
    add_textbox(
        slide,
        Inches(0.7),
        Inches(1.35),
        Inches(12),
        Inches(0.45),
        "※画像プロンプトに詳細が薄いため、Criticの限界を踏まえた自己検証型RAGとして補完（仮定を明示）",
        size=12,
        color=ORANGE,
    )
    add_bullets(
        slide,
        Inches(0.7),
        Inches(1.9),
        Inches(12),
        Inches(1.5),
        [
            "Criticは「生成後の評価」に強いが、検索不足・根拠不足を途中で止められない",
            "Self-RAGでは、モデル自身が「検索が必要か」「根拠は十分か」を判定する",
            "不足なら再検索・再生成し、品質ゲートを通過したものだけを返す",
        ],
        size=16,
    )
    flow_boxes(
        slide,
        ["入力", "要検索判定", "検索/生成", "自己評価", "出力 or 再実行"],
        top=Inches(4.0),
        box_w=Inches(2.1),
        box_h=Inches(0.85),
        gap=Inches(0.25),
    )
    add_textbox(
        slide,
        Inches(0.7),
        Inches(5.3),
        Inches(12),
        Inches(1.2),
        "判定例：検索不要 / 再検索が必要 / 根拠不足で再生成 / 品質OKで確定",
        size=15,
        color=GRAY,
    )
    add_footer(slide, 10)
    return "Self-RAG"


def slide_11_self_rag_improve(prs):
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    add_header_bar(slide, "Self-RAG改善", "改善内容と結果")
    items = [
        ("改善1", "検索要否の自己チェックを導入し、無駄な検索と検索漏れを削減"),
        ("改善2", "回答根拠（参照障害票）の十分性をチェックし、根拠不足なら再検索"),
        ("改善3", "Critic評価と組み合わせ、最終出力の安定性を向上"),
    ]
    for i, (h, body) in enumerate(items):
        top = Inches(1.4) + i * Inches(1.5)
        add_rounded_box(slide, Inches(0.5), top, Inches(1.8), Inches(1.1), ACCENT, h, size=16)
        add_rect(slide, Inches(2.5), top, Inches(10.3), Inches(1.1), LIGHT_BG)
        add_textbox(slide, Inches(2.8), top + Inches(0.3), Inches(9.8), Inches(0.6), body, size=16, color=DARK)
    add_footer(slide, 11)
    return "Self-RAG改善"


def slide_12_evolution(prs):
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    add_header_bar(slide, "試行錯誤の変遷", "課題と打ち手の俯瞰")

    rows = [
        ("Basic RAG", "表現ゆれに弱い", "検索品質が回答品質を決める"),
        ("Query Rewrite", "変換しすぎて意味が消える", "原文維持・JSON制約"),
        ("Critic RAG", "無関係テーマが高評価", "ユーザー質問を評価に含める"),
        ("Self-RAG", "途中失敗を止められない", "自己判定＋再検索ループ"),
    ]
    # header
    headers = ["段階", "発生した課題", "改善の打ち手"]
    widths = [Inches(2.8), Inches(4.6), Inches(4.6)]
    lefts = [Inches(0.5), Inches(3.3), Inches(7.9)]
    for left, w, h in zip(lefts, widths, headers):
        add_rounded_box(slide, left, Inches(1.4), w, Inches(0.55), NAVY, h, size=14)

    for i, (stage, issue, fix) in enumerate(rows):
        top = Inches(2.15) + i * Inches(1.05)
        bg = LIGHT_BG if i % 2 == 0 else WHITE
        add_rect(slide, Inches(0.5), top, Inches(12.3), Inches(0.95), bg, line=RGBColor(0xDD, 0xDD, 0xDD))
        add_textbox(slide, lefts[0] + Inches(0.15), top + Inches(0.25), widths[0] - Inches(0.2), Inches(0.5), stage, size=15, bold=True, color=NAVY)
        add_textbox(slide, lefts[1] + Inches(0.15), top + Inches(0.25), widths[1] - Inches(0.2), Inches(0.5), issue, size=14, color=DARK)
        add_textbox(slide, lefts[2] + Inches(0.15), top + Inches(0.25), widths[2] - Inches(0.2), Inches(0.5), fix, size=14, color=DARK)
    add_footer(slide, 12)
    return "試行錯誤の変遷"


def slide_13_evaluation(prs):
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    add_header_bar(slide, "評価結果", "定性評価（数値は例示）")
    add_textbox(
        slide,
        Inches(0.7),
        Inches(1.3),
        Inches(12),
        Inches(0.35),
        "※定量値は画像プロンプトに無いため、PoCで想定される相対改善の例示です",
        size=12,
        color=ORANGE,
    )

    headers = ["観点", "Basic", "Rewrite", "Critic", "Self-RAG"]
    data = [
        ["表現ゆれ耐性", "△", "○", "○", "◎"],
        ["検索ヒット率（例示）", "低", "中", "中", "高"],
        ["テーマ妥当性", "△", "△", "○", "◎"],
        ["無関係提案の抑制", "×", "△", "○", "◎"],
        ["出力安定性", "△", "○", "○", "◎"],
    ]
    col_w = [Inches(3.2), Inches(2.1), Inches(2.1), Inches(2.1), Inches(2.1)]
    left0 = Inches(0.6)
    # header row
    x = left0
    for i, h in enumerate(headers):
        add_rounded_box(slide, x, Inches(1.85), col_w[i] - Inches(0.1), Inches(0.55), NAVY, h, size=13)
        x += col_w[i]
    for r, row in enumerate(data):
        x = left0
        top = Inches(2.55) + r * Inches(0.75)
        for c, cell in enumerate(row):
            fill = LIGHT_BG if r % 2 == 0 else WHITE
            add_rect(slide, x, top, col_w[c] - Inches(0.1), Inches(0.65), fill, line=RGBColor(0xDD, 0xDD, 0xDD))
            color = NAVY if c == 0 else DARK
            bold = c == 0
            add_textbox(slide, x + Inches(0.1), top + Inches(0.12), col_w[c] - Inches(0.3), Inches(0.45), cell, size=14, bold=bold, color=color, align=PP_ALIGN.CENTER if c else PP_ALIGN.LEFT)
            x += col_w[c]
    add_footer(slide, 13)
    return "評価結果"


def slide_14_insights(prs):
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    add_header_bar(slide, "PoCで得られた知見", "再現可能な学び")
    insights = [
        ("1", "回答品質は検索品質に依存する", "生成プロンプト改善だけでは限界。まず当てる検索を作る"),
        ("2", "クエリ変換は制約設計が命", "「検索しやすく」だけでは過剰一般化が起きる"),
        ("3", "評価にはユーザー質問を必ず含める", "テーマ単体評価は誤高評価を招く"),
        ("4", "段階的にゲートを増やす", "Rewrite → Critic → Self で失敗モードを潰す"),
    ]
    for i, (n, title, body) in enumerate(insights):
        top = Inches(1.35) + i * Inches(1.25)
        add_rounded_box(slide, Inches(0.5), top, Inches(0.8), Inches(1.0), ACCENT, n, size=20)
        add_textbox(slide, Inches(1.5), top + Inches(0.1), Inches(11), Inches(0.4), title, size=17, bold=True, color=NAVY)
        add_textbox(slide, Inches(1.5), top + Inches(0.5), Inches(11), Inches(0.4), body, size=14, color=GRAY)
    add_footer(slide, 14)
    return "PoCで得られた知見"


def slide_15_future(prs):
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    add_header_bar(slide, "今後の展望", "PoCから次のステップへ")
    items = [
        ("ナレッジ拡充", "本番障害票の追加投入とチャンク設計の見直し"),
        ("評価自動化", "関連性・再発防止・実現可能性の自動スコアリング定着"),
        ("運用組み込み", "改善テーマ提案を運用改善会議のインプットにする"),
        ("継続学習", "採用された改善テーマをナレッジへフィードバック"),
    ]
    for i, (t, d) in enumerate(items):
        left = Inches(0.5) + (i % 2) * Inches(6.4)
        top = Inches(1.5) + (i // 2) * Inches(2.4)
        add_rect(slide, left, top, Inches(6.1), Inches(2.1), LIGHT_BG)
        add_rounded_box(slide, left + Inches(0.25), top + Inches(0.3), Inches(5.6), Inches(0.55), NAVY, t, size=16)
        add_textbox(slide, left + Inches(0.35), top + Inches(1.1), Inches(5.4), Inches(0.7), d, size=15, color=DARK)
    add_footer(slide, 15)
    return "今後の展望"


def slide_16_summary(prs):
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    add_header_bar(slide, "まとめ", "課題→改善の繰り返しで品質を段階的に向上")
    add_bullets(
        slide,
        Inches(0.7),
        Inches(1.5),
        Inches(12),
        Inches(2.5),
        [
            "障害票ナレッジを使い、類似検索・情報抽出・改善テーマ生成をDify上で実現",
            "Basic RAGの表現ゆれ課題を、Query Rewriteの制約設計で改善",
            "Criticでテーマ品質を評価し、ユーザー質問を評価対象に含めることで誤高評価を抑制",
            "Self-RAGで検索・根拠の自己判定と再実行を加え、出力を安定化",
        ],
        size=17,
    )
    add_rect(slide, Inches(0.5), Inches(4.5), Inches(12.3), Inches(2.0), NAVY)
    add_textbox(
        slide,
        Inches(0.9),
        Inches(5.0),
        Inches(11.5),
        Inches(1.0),
        "単なるシステム紹介ではなく、\n「発生した課題 → 原因分析 → 改善 → 結果」の連続が本PoCの本質",
        size=18,
        bold=True,
        color=WHITE,
        align=PP_ALIGN.CENTER,
    )
    add_footer(slide, 16)
    return "まとめ"


BUILDERS = [
    slide_01_title,
    slide_02_background,
    slide_03_architecture,
    slide_04_basic_rag,
    slide_05_basic_issues,
    slide_06_query_rewrite,
    slide_07_query_rewrite_result,
    slide_08_critic,
    slide_09_critic_improve,
    slide_10_self_rag,
    slide_11_self_rag_improve,
    slide_12_evolution,
    slide_13_evaluation,
    slide_14_insights,
    slide_15_future,
    slide_16_summary,
]


# ---------------------------------------------------------------------------
# HTML preview (same narrative content)
# ---------------------------------------------------------------------------

HTML_SLIDES = [
    {
        "title": "タイトル",
        "body": """
        <p class="eyebrow">AI技術検証プロジェクト｜PoC報告</p>
        <h1>障害票を活用した<br>運用改善テーマ発見AIのPoC検証</h1>
        <p class="sub">Basic RAG → Query Rewrite → Critic → Self-RAG による段階的改善</p>
        """,
    },
    {
        "title": "背景・目的",
        "body": """
        <ul>
          <li>障害票をナレッジとして蓄積し、運用改善テーマを自動生成するAIを構築した</li>
          <li>本発表では「システム紹介」ではなく、構築〜改善の試行錯誤プロセスを説明する</li>
          <li>各段階を「発生した課題 → 原因分析 → 改善 → 結果」の流れで整理する</li>
        </ul>
        <div class="flow"><span>発生した課題</span><span>→</span><span>原因分析</span><span>→</span><span>改善</span><span>→</span><span>結果</span></div>
        """,
    },
    {
        "title": "システム全体構成",
        "body": """
        <ul>
          <li>障害票をナレッジベースとして登録（Dify）</li>
          <li>ユーザー入力に対し、類似障害検索 / 障害情報抽出 / 改善テーマ生成を実行</li>
        </ul>
        <div class="cards">
          <div class="card"><h3>類似障害検索</h3><p>過去の類似チケットを検索</p></div>
          <div class="card"><h3>障害情報抽出</h3><p>原因・対応・恒久対策を整理</p></div>
          <div class="card"><h3>改善テーマ生成</h3><p>テーマ・内容・期待効果を提案</p></div>
        </div>
        <div class="flow"><span>Basic RAG</span><span>→</span><span>Query Rewrite</span><span>→</span><span>Critic RAG</span><span>→</span><span>Self-RAG</span></div>
        """,
    },
    {
        "title": "Basic RAG",
        "body": """
        <div class="flow"><span>入力</span><span>→</span><span>ナレッジ検索</span><span>→</span><span>障害情報抽出</span><span>→</span><span>改善テーマ生成</span></div>
        <div class="cards">
          <div class="card"><h3>【障害情報抽出】</h3><p>あなたは障害分析担当です。<br>検索結果から 障害名 / 原因 / 対応 / 恒久対策 を整理</p></div>
          <div class="card"><h3>【改善テーマ生成】</h3><p>あなたは運用改善担当です。<br>改善テーマ / 改善内容 / 期待効果 を作成</p></div>
        </div>
        """,
    },
    {
        "title": "Basic RAGで発生した課題",
        "body": """
        <div class="cards">
          <div class="card good"><h3>成功例</h3><ul><li>バックアップ失敗</li><li>DB接続エラー</li><li>容量不足による障害</li></ul></div>
          <div class="card bad"><h3>課題：表現ゆれに弱い</h3><ul><li>DBトラブル</li><li>バックアップ問題</li><li>言い回しが違うと検索が外れる</li></ul></div>
          <div class="card"><h3>学び</h3><ul><li>回答品質は検索品質に依存する</li><li>生成改善だけでは限界</li></ul></div>
        </div>
        """,
    },
    {
        "title": "Query Rewrite RAG",
        "body": """
        <p>目的：表現ゆれ対策。検索前にクエリを書き換える。</p>
        <div class="card bad"><h3>失敗したプロンプト</h3><p>「質問を検索しやすい形式へ変換してください」</p></div>
        <div class="card bad"><h3>問題：変換しすぎた</h3><p>例）DBトラブル → 運用改善課題<br>具体的な障害語が消え、検索精度が悪化</p></div>
        """,
    },
    {
        "title": "Query Rewriteによる改善結果",
        "body": """
        <div class="cards">
          <div class="card"><h3>原文を維持</h3></div>
          <div class="card"><h3>説明しない</h3></div>
          <div class="card"><h3>言い換えすぎない</h3></div>
          <div class="card"><h3>JSONのみ出力</h3></div>
        </div>
        <div class="cards">
          <div class="card good"><h3>DBトラブル</h3><p>→ 検索成功</p></div>
          <div class="card good"><h3>バックアップ問題</h3><p>→ 検索成功</p></div>
          <div class="card good"><h3>容量管理不足</h3><p>→ 検索成功</p></div>
        </div>
        """,
    },
    {
        "title": "Critic RAG",
        "body": """
        <p>目的：改善テーマの品質評価</p>
        <div class="card bad"><h3>初期プロンプト</h3><p>あなたは運用改善レビュー担当です。<br>改善テーマの妥当性を評価してください。</p></div>
        <div class="cards">
          <div class="card"><h3>関連性</h3></div>
          <div class="card"><h3>再発防止効果</h3></div>
          <div class="card"><h3>実現可能性</h3></div>
        </div>
        """,
    },
    {
        "title": "Critic改善",
        "body": """
        <div class="cards">
          <div class="card bad"><h3>問題</h3><p>無関係テーマでも総合評価◎</p></div>
          <div class="card warn"><h3>原因</h3><p>ユーザー質問を評価対象に含めていなかった</p></div>
          <div class="card good"><h3>改善</h3><p>評価入力にユーザー質問を追加し、関連性を必須ゲート化</p></div>
        </div>
        """,
    },
    {
        "title": "Self-RAG",
        "body": """
        <p class="note">※画像プロンプトに詳細が薄いため仮定を明示して補完</p>
        <ul>
          <li>検索要否・根拠十分性を自己判定</li>
          <li>不足なら再検索・再生成</li>
          <li>品質ゲート通過のみ出力</li>
        </ul>
        <div class="flow"><span>入力</span><span>→</span><span>要検索判定</span><span>→</span><span>検索/生成</span><span>→</span><span>自己評価</span><span>→</span><span>出力/再実行</span></div>
        """,
    },
    {
        "title": "Self-RAG改善",
        "body": """
        <ul>
          <li>検索要否の自己チェックを導入</li>
          <li>回答根拠（参照障害票）の十分性チェック</li>
          <li>Critic評価と組み合わせて出力安定性を向上</li>
        </ul>
        """,
    },
    {
        "title": "試行錯誤の変遷",
        "body": """
        <table>
          <tr><th>段階</th><th>発生した課題</th><th>改善の打ち手</th></tr>
          <tr><td>Basic RAG</td><td>表現ゆれに弱い</td><td>検索品質が回答品質を決める</td></tr>
          <tr><td>Query Rewrite</td><td>変換しすぎて意味が消える</td><td>原文維持・JSON制約</td></tr>
          <tr><td>Critic RAG</td><td>無関係テーマが高評価</td><td>ユーザー質問を評価に含める</td></tr>
          <tr><td>Self-RAG</td><td>途中失敗を止められない</td><td>自己判定＋再検索ループ</td></tr>
        </table>
        """,
    },
    {
        "title": "評価結果",
        "body": """
        <p class="note">※定量値は例示</p>
        <table>
          <tr><th>観点</th><th>Basic</th><th>Rewrite</th><th>Critic</th><th>Self-RAG</th></tr>
          <tr><td>表現ゆれ耐性</td><td>△</td><td>○</td><td>○</td><td>◎</td></tr>
          <tr><td>検索ヒット率（例示）</td><td>低</td><td>中</td><td>中</td><td>高</td></tr>
          <tr><td>テーマ妥当性</td><td>△</td><td>△</td><td>○</td><td>◎</td></tr>
          <tr><td>無関係提案の抑制</td><td>×</td><td>△</td><td>○</td><td>◎</td></tr>
          <tr><td>出力安定性</td><td>△</td><td>○</td><td>○</td><td>◎</td></tr>
        </table>
        """,
    },
    {
        "title": "PoCで得られた知見",
        "body": """
        <ol>
          <li><strong>回答品質は検索品質に依存する</strong></li>
          <li><strong>クエリ変換は制約設計が命</strong></li>
          <li><strong>評価にはユーザー質問を必ず含める</strong></li>
          <li><strong>段階的にゲートを増やす</strong></li>
        </ol>
        """,
    },
    {
        "title": "今後の展望",
        "body": """
        <div class="cards">
          <div class="card"><h3>ナレッジ拡充</h3><p>本番障害票の追加とチャンク設計見直し</p></div>
          <div class="card"><h3>評価自動化</h3><p>3軸スコアリングの定着</p></div>
          <div class="card"><h3>運用組み込み</h3><p>改善会議のインプット化</p></div>
          <div class="card"><h3>継続学習</h3><p>採用テーマのフィードバック</p></div>
        </div>
        """,
    },
    {
        "title": "まとめ",
        "body": """
        <ul>
          <li>障害票ナレッジで類似検索・抽出・テーマ生成を実現（Dify）</li>
          <li>表現ゆれは Query Rewrite の制約設計で改善</li>
          <li>Critic でテーマ品質を担保（ユーザー質問を評価に含める）</li>
          <li>Self-RAG で自己判定と再実行を追加し安定化</li>
        </ul>
        <div class="hero-note">課題 → 原因分析 → 改善 → 結果 の連続が本PoCの本質</div>
        """,
    },
]


def write_html(titles: list[str]):
    sections = []
    for i, (title, slide) in enumerate(zip(titles, HTML_SLIDES), start=1):
        sections.append(
            f'<section class="slide" id="s{i}">'
            f'<header><span class="num">{i} / 16</span><h2>{slide["title"]}</h2></header>'
            f'<div class="content">{slide["body"]}</div>'
            f"</section>"
        )
    html = f"""<!DOCTYPE html>
<html lang="ja">
<head>
<meta charset="utf-8"/>
<meta name="viewport" content="width=device-width, initial-scale=1"/>
<title>障害票RAG PoC 発表資料（プレビュー）</title>
<style>
  :root {{
    --navy: #1b3a5f;
    --accent: #2e6b9e;
    --bg: #f5f7fa;
    --dark: #2c2c2c;
    --gray: #5a5a5a;
    --good: #2d6a4f;
    --bad: #a63d40;
    --warn: #c46b2e;
  }}
  * {{ box-sizing: border-box; }}
  body {{
    margin: 0;
    font-family: "Yu Gothic", "Hiragino Sans", "Noto Sans CJK JP", "WenQuanYi Micro Hei", sans-serif;
    background: #e8eef4;
    color: var(--dark);
  }}
  .toolbar {{
    position: sticky; top: 0; z-index: 10;
    display: flex; gap: 8px; flex-wrap: wrap; align-items: center;
    padding: 10px 16px; background: var(--navy); color: #fff;
  }}
  .toolbar button {{
    background: var(--accent); color: #fff; border: 0; border-radius: 6px;
    padding: 8px 12px; cursor: pointer; font-size: 14px;
  }}
  .toolbar .title {{ margin-right: auto; font-weight: 700; }}
  main {{ max-width: 1100px; margin: 24px auto; padding: 0 16px 48px; }}
  .slide {{
    background: #fff; border-radius: 12px; padding: 28px 32px 36px;
    margin-bottom: 28px; box-shadow: 0 8px 24px rgba(27,58,95,.08);
    border-top: 6px solid var(--navy);
  }}
  .slide header {{ display: flex; align-items: baseline; gap: 16px; margin-bottom: 16px; }}
  .slide h2 {{ margin: 0; color: var(--navy); font-size: 28px; }}
  .num {{ color: var(--accent); font-weight: 700; }}
  ul, ol {{ line-height: 1.7; font-size: 17px; }}
  .flow {{
    display: flex; flex-wrap: wrap; gap: 8px; align-items: center;
    margin: 18px 0; font-weight: 700; color: var(--navy);
  }}
  .flow span {{ background: var(--bg); padding: 8px 12px; border-radius: 8px; }}
  .cards {{ display: grid; grid-template-columns: repeat(auto-fit, minmax(220px, 1fr)); gap: 12px; margin-top: 16px; }}
  .card {{ background: var(--bg); border-radius: 10px; padding: 14px 16px; }}
  .card h3 {{ margin: 0 0 8px; color: var(--navy); font-size: 16px; }}
  .card.good {{ border-left: 4px solid var(--good); }}
  .card.bad {{ border-left: 4px solid var(--bad); }}
  .card.warn {{ border-left: 4px solid var(--warn); }}
  table {{ width: 100%; border-collapse: collapse; margin-top: 12px; font-size: 15px; }}
  th, td {{ border: 1px solid #ddd; padding: 10px; text-align: left; }}
  th {{ background: var(--navy); color: #fff; }}
  tr:nth-child(even) td {{ background: var(--bg); }}
  .note {{ color: var(--warn); font-size: 13px; }}
  .eyebrow {{ color: var(--accent); font-weight: 700; }}
  h1 {{ color: var(--navy); font-size: 34px; line-height: 1.35; }}
  .sub {{ color: var(--gray); }}
  .hero-note {{
    margin-top: 20px; background: var(--navy); color: #fff;
    padding: 18px 20px; border-radius: 10px; font-weight: 700; text-align: center;
  }}
</style>
</head>
<body>
  <div class="toolbar">
    <div class="title">障害票RAG PoC 発表資料（HTMLプレビュー）</div>
    <button onclick="location.hash='s1'">先頭</button>
    <button onclick="step(-1)">前へ</button>
    <button onclick="step(1)">次へ</button>
  </div>
  <main>
    {''.join(sections)}
  </main>
  <script>
    let idx = 1;
    function step(d) {{
      idx = Math.min(16, Math.max(1, idx + d));
      location.hash = 's' + idx;
      document.getElementById('s' + idx)?.scrollIntoView({{behavior:'smooth'}});
    }}
    window.addEventListener('hashchange', () => {{
      const n = Number((location.hash || '#s1').slice(2));
      if (n) idx = n;
    }});
  </script>
</body>
</html>
"""
    HTML_PATH.write_text(html, encoding="utf-8")


def verify_pptx(titles: list[str]):
    assert len(titles) == 16, f"expected 16 slides, got {len(titles)}"
    assert titles == EXPECTED_TITLES, f"title mismatch:\n{titles}\n!=\n{EXPECTED_TITLES}"
    prs = Presentation(str(PPTX_PATH))
    assert len(prs.slides) == 16, f"pptx slide count {len(prs.slides)}"
    print("VERIFY OK: 16 slides, titles match")


def main():
    prs = Presentation()
    prs.slide_width = SLIDE_W
    prs.slide_height = SLIDE_H
    # remove default empty layout dependency by using blank layout index 6 if present;
    # python-pptx default template has layout 6 as blank.
    titles = []
    for builder in BUILDERS:
        titles.append(builder(prs))

    prs.save(str(PPTX_PATH))
    write_html(titles)
    print(f"Wrote {PPTX_PATH}")
    print(f"Wrote {HTML_PATH}")
    verify_pptx(titles)
    for i, t in enumerate(titles, 1):
        print(f"  {i:2d}. {t}")


if __name__ == "__main__":
    main()
